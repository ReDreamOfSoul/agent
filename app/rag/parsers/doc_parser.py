# -*- coding: utf-8 -*-
"""DOC/DOCX 解析器

策略链（按顺序尝试）：
1. .docx → python-docx 直接解析；
2. .doc（老格式）→ LibreOffice 无头转换（若可用）→ Windows Word COM（若可用）→ 报错并给出提示。
题库、教学大纲等课程资料多为 .doc 老格式，本模块保证其可入库。
"""
from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from .pdf_parser import TextChunk, _split_with_overlap

logger = logging.getLogger("agent.parser.doc")


def _parse_docx_text(path: Path) -> str:
    """python-docx 提取段落与表格文本"""
    import docx
    d = docx.Document(str(path))
    parts: list[str] = []
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _soffice_to_txt(path: Path) -> str | None:
    """尝试用 LibreOffice 无头转换 .doc → txt"""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        out_dir = Path(tmp)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "txt:Text (encoded):UTF8",
                 str(path), "--outdir", str(out_dir)],
                check=True, capture_output=True, timeout=120,
            )
            out_file = out_dir / (path.stem + ".txt")
            if out_file.exists():
                return out_file.read_text(encoding="utf-8", errors="ignore")
        except Exception as exc:  # noqa: BLE001
            logger.warning("LibreOffice 转换 %s 失败: %s", path.name, exc)
    return None


def _word_com_to_txt(path: Path) -> str | None:
    """Windows 下用 Word COM 提取 .doc 文本"""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(str(path.resolve()), ReadOnly=True)
            try:
                text = doc.Content.Text
            finally:
                doc.Close(False)
            return text
        finally:
            word.Quit()
    except Exception as exc:  # noqa: BLE001
        logger.warning("Word COM 提取 %s 失败: %s", path.name, exc)
    return None


def parse_doc(path: str | Path, chunk_size: int = 500, chunk_overlap: int = 60) -> list[TextChunk]:
    """解析 DOC/DOCX 为文本块"""
    path = Path(path)
    suffix = path.suffix.lower()
    text = ""
    if suffix == ".docx":
        text = _parse_docx_text(path)
    elif suffix == ".doc":
        text = _soffice_to_txt(path) or _word_com_to_txt(path) or ""
    else:
        raise ValueError(f"不支持的文档类型: {suffix}")

    if not text.strip():
        raise RuntimeError(
            f"文档 {path.name} 文本提取失败。.doc 老格式文件请安装 LibreOffice 或确保本机安装 "
            "Microsoft Word（Windows 环境自动使用 Word COM 提取）。"
        )
    chunks = [
        TextChunk(text=c, source=path.name, chapter="", page=0, resource_type="题库/教学大纲")
        for c in _split_with_overlap(text, chunk_size, chunk_overlap)
    ]
    logger.info("DOC 解析完成: %s → %d 块", path.name, len(chunks))
    return chunks
